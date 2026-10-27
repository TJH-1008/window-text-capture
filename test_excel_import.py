import json
import tempfile
import unittest
import zipfile
from pathlib import Path

import excel_import as importer


def shop(name="测试店", floor="商场1F", category="咖啡"):
    return f"{name}\n★\n★\n4.5\n12条\n￥28/人\n{category}\n八宝街\n{floor}\n"


def snapshot(text, extra=None):
    docs = [{"name": "历史", "text": shop("旧页面店")},
            {"name": "Page-Frame", "text": "输入商户名、地点、找优惠\n全部楼层\n全部分类\n" + text}]
    if extra:
        docs.insert(1, {"name": "Identity", "text": extra})
    return {"documents": docs, "primary_document_index": len(docs) - 1,
            "captured_at": "2026/10/21 15:13:25"}


class ExcelImportTests(unittest.TestCase):
    def test_floor_label_variants_preserve_shops_and_original_floor(self):
        labels = ["L1", "L2", "L3", "l4", "南馆L2", "南馆B1", "A区3L",
                  "1期L2", "东区2F", "北区B1", "F1", "F2", "1L", "2L",
                  "G层", "LG层", "A区G层", "P1", "A区P1", "L1,L2", "L5,L6"]
        for label in labels:
            with self.subTest(label=label):
                raw = shop("楼层测试店", "商场" + label)
                rows, issues = importer.parse_list(snapshot(raw)["documents"][-1]["text"])
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["name"], "楼层测试店")
                self.assertEqual(rows[0]["floor"], label.upper())
                self.assertEqual(rows[0]["floor_raw"], "商场" + label)
                self.assertEqual(rows[0]["category"], "咖啡")
                self.assertEqual(issues, [])

    def test_floor_label_truncation_stays_visible_and_unknown_is_not_guessed(self):
        for raw, expected in [("商场L5商…", "L5"), ("商场北区B1商场北…", "北区B1")]:
            with self.subTest(raw=raw):
                rows, issues = importer.parse_list(shop(floor=raw))
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["floor"], expected)
                self.assertEqual(rows[0]["floor_raw"], raw)
                self.assertTrue(any("截断" in i for i in rows[0]["issues"]))
        for label in ("商场XYZ", "商场L", "商场服务", "商场优惠L2", "商场未知L2"):
            with self.subTest(label=label):
                self.assertEqual(importer.parse_list(shop(floor=label))[0], [])

    def test_l_floors_round_trip_excel_without_losing_b1_or_other_shops(self):
        import openpyxl
        floors = ["B1", "L2", "L3", "南馆L2", "A区3L", "F2", "LG层", "L5,L6"]
        text = "".join(shop("店" + str(i), "商场" + floor) for i, floor in enumerate(floors))
        result = importer.convert_session("楼层回归商场", [snapshot(text)])
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "result.xlsx"
            importer.write_result(path, result)
            book = openpyxl.load_workbook(path)
            values = list(book["品牌明细"].iter_rows(min_row=2, values_only=True))
            self.assertEqual([row[3] for row in values], ["店" + str(i) for i in range(len(floors))])
            self.assertEqual([row[4] for row in values], floors)
            self.assertEqual([row[7] for row in values], ["咖啡"] * len(floors))
            book.close()

    def test_fast_capture_requires_bottom_stability_for_export_status(self):
        snap = snapshot(shop() + "已经到底了")
        snap.update(collection_mode="final_only", bottom_stable=False, bottom_confirmed=False)
        result = importer.convert_session("测试商场", [snap])
        self.assertFalse(result["bottom_confirmed"])
        self.assertTrue(any("底部" in issue for issue in result["issues"]))
        snap.update(bottom_stable=True, bottom_confirmed=True)
        self.assertTrue(importer.convert_session("测试商场", [snap])["bottom_confirmed"])

    def test_folder_names_override_capture_labels_in_export_and_workbook(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / 'captures'
            for name in ('茂业天地', '20261022_101657_372497_群光广场', '20261022_111657_372498_群光广场'):
                folder = source / name
                folder.mkdir(parents=True)
                (folder / 'capture.json').write_text(json.dumps(dict(snapshot(shop()), merchant_name='大众点评美食电影运动旅游门票')), encoding='utf-8')
            result = importer.batch_export([source], root / 'out')
            self.assertEqual(result['success_count'], 3)
            labels = []
            for output in result['outputs']:
                book = openpyxl.load_workbook(output, read_only=True)
                try:
                    label = list(book['品牌明细'].values)[1][2]
                    self.assertTrue(Path(output).name.endswith(label + '_品牌业态明细.xlsx'))
                    labels.append(label)
                finally:
                    book.close()
            self.assertCountEqual(labels, ['茂业天地', '群光广场', '群光广场'])
            self.assertEqual(len(set(result['outputs'])), 3)

    def test_zip_uses_parent_folder_name_and_archive_stem_for_root_capture(self):
        with tempfile.TemporaryDirectory() as temp:
            archive = Path(temp) / '单家商场.zip'
            with zipfile.ZipFile(archive, 'w') as z:
                data = json.dumps(dict(snapshot(shop()), merchant_name='错误窗口标题'))
                z.writestr('20261022_101657_372497_仁和春天百货/capture.json', data)
                z.writestr('capture.json', data)
            sessions = list(importer.zip_sessions(archive))
            self.assertCountEqual([item[0] for item in sessions], ['仁和春天百货', '单家商场'])

    def test_text_import_also_uses_containing_folder(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / 'COSMO'
            folder.mkdir()
            path = folder / '任意文件名.txt'
            path.write_text('输入商户名\n全部楼层\n全部分类\n' + shop(), encoding='utf-8')
            self.assertEqual(importer.read_session(path)[0], 'COSMO')

    def test_parser_uses_floor_boundary_and_preserves_unknowns(self):
        rows, issues = importer.parse_list(snapshot(shop() + shop("其他店", "商场内", "新业态"))["documents"][-1]["text"])
        self.assertEqual([r["name"] for r in rows], ["测试店", "其他店"])
        self.assertEqual(rows[0]["group"], "餐饮")
        self.assertEqual(rows[1]["group"], "其他/待复核")
        self.assertIn("未注明具体楼层", rows[1]["issues"])
        self.assertEqual(issues, [])

    def test_snapshot_union_preserves_different_floors_and_same_snapshot_duplicates(self):
        first = snapshot(shop("A", "商场1F") * 2)
        second = snapshot(shop("A", "商场1F") + shop("A", "商场2F"))
        result = importer.convert_session("测试商场", [first, second])
        self.assertEqual(len(result["records"]), 3)
        self.assertNotIn("旧页面店", [r["name"] for r in result["records"]])
        self.assertTrue(any("重复" in "；".join(r["issues"]) for r in result["records"]))

    def test_header_without_readable_rows_remains_review_not_success(self):
        result = importer.convert_session("测试商场", [snapshot("未知内容\n商场XYZ\n已经到底了")])
        self.assertEqual(result["records"], [])
        self.assertTrue(result["issues"])

    def test_current_non_list_page_does_not_parse_historical_list(self):
        snap = snapshot(shop("历史列表商户"))
        snap["documents"].append({"name": "Current", "text": "商场首页\n商场服务"})
        snap["primary_document_index"] = len(snap["documents"]) - 1
        result = importer.convert_session("测试商场", [snap])
        self.assertEqual(result["records"], [])
        self.assertTrue(any("当前文档" in issue for issue in result["issues"]))

    def test_incomplete_filters_and_count_difference_are_explicit(self):
        snap = snapshot(shop(), "测试商场\n商场服务\n全部\n(3)")
        snap["documents"][-1]["text"] = snap["documents"][-1]["text"].replace("全部分类", "餐饮")
        result = importer.convert_session("测试商场", [snap])
        self.assertEqual(result["platform_count"], 3)
        self.assertTrue(any("分类" in x for x in result["issues"]))
        self.assertTrue(any("数量" in x for x in result["issues"]))

    def test_batch_zip_import_does_not_extract_and_keeps_unreadable_input(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "batch.zip"
            with zipfile.ZipFile(archive, "w") as z:
                z.writestr("A/capture.json", json.dumps(dict(snapshot(shop()), merchant_name="A")))
                z.writestr("B/capture.json", "broken json")
                z.writestr("../never-write.txt", "outside")
            result = importer.batch_export([archive], root / "out")
            self.assertEqual(result["success_count"], 1)
            self.assertEqual(result["error_count"], 1)
            self.assertFalse((root.parent / "never-write.txt").exists())
            self.assertTrue(Path(result["index"]).exists())

    def test_excel_headers_original_text_and_formula_like_name_round_trip(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "result.xlsx"
            value = '=HYPERLINK("x")'
            result = importer.convert_session("测试商场", [snapshot(shop(value))])
            importer.write_result(path, result)
            book = openpyxl.load_workbook(path)
            sheet = book["品牌明细"]
            self.assertEqual([c.value for c in sheet[1]], importer.HEADERS)
            self.assertEqual(sheet["D2"].value, value)
            self.assertEqual(sheet["D2"].data_type, "s")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertIn("原始文本", book.sheetnames)
            self.assertIn("待复核记录", book.sheetnames)
            book.close()
            book = openpyxl.load_workbook(path, read_only=True)
            try:
                self.assertEqual(book["品牌明细"].max_column, 11)
                self.assertEqual(len(list(book["品牌明细"].values)[1]), 11)
            finally:
                book.close()
            with self.assertRaises(FileExistsError):
                importer.write_result(path, result)

    def test_template_order_and_custom_mapping(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = root / "template.xlsx"
            headers = list(reversed(importer.HEADERS)) + ["自定义字段"]
            importer.write_xlsx(template, [("品牌明细", [headers], [20] * len(headers))])
            self.assertEqual(importer.read_template(template), headers)
            result = importer.convert_session("测试商场", [snapshot(shop(category="新业态"))], {"新业态": "零售"})
            self.assertEqual(result["records"][0]["group"], "零售")

    def test_original_brand_detail_sheet_name_is_accepted(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'original.xlsx'
            importer.write_xlsx(path, [('品牌明细表', [importer.HEADERS], [20] * 11)])
            self.assertEqual(importer.read_template(path), importer.HEADERS)

    def test_write_xlsx_splits_oversized_cells_into_continuation_sheet(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'long.xlsx'
            original = '王府井问题；' * 12000
            importer.write_xlsx(path, [('批次索引', [['商场', '待核事项'], ['王府井', original]], [30, 100])])
            book = openpyxl.load_workbook(path, read_only=True)
            try:
                sheet = book['批次索引']
                preview = sheet['B2'].value
                self.assertIn('超长文本续表', preview)
                continuation = book['超长文本续表']
                rows = list(continuation.values)
                self.assertEqual(rows[0], ('来源工作表', '单元格', '分段', '完整内容'))
                restored = ''.join(row[3] for row in rows[1:] if row[1] == 'B2')
                self.assertEqual(restored, original)
                self.assertTrue(all(len(row[3].encode('utf-16-le')) // 2 <= 32000 for row in rows[1:]))
            finally:
                book.close()

    def test_utf16_overflow_does_not_overwrite_existing_sheet_or_formula_like_text(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'unicode.xlsx'
            original = '=' + '\U0001f600' * 18000
            boundary = 'A' * 32767
            importer.write_xlsx(path, [('数据', [['长文', '边界'], [original, boundary]], [40, 40]),
                                       ('超长文本续表', [['原有工作表'], ['保留']], [30])])
            book = openpyxl.load_workbook(path, read_only=True)
            try:
                self.assertEqual(book['数据']['B2'].value, boundary)
                self.assertEqual(book['数据']['A2'].data_type, 's')
                self.assertEqual(book['超长文本续表']['A2'].value, '保留')
                self.assertIn('超长文本续表_2', book['数据']['A2'].value)
                chunks = list(book['超长文本续表_2'].values)[1:]
                self.assertEqual(''.join(row[3] for row in chunks), original)
                self.assertTrue(all(len(row[3].encode('utf-16-le')) // 2 <= 32767 for row in chunks))
            finally:
                book.close()

    def test_batch_index_handles_many_unrecognized_floor_warnings(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / '王府井' / 'capture.json'
            source.parent.mkdir()
            text = shop() + ''.join('商场无法识别的标签' + str(i) + '\n' for i in range(1800))
            snap = snapshot(text)
            source.write_text(json.dumps(snap), encoding='utf-8')
            expected = '；'.join(importer.convert_session('王府井', [snap], source=str(source))['issues'])
            self.assertGreater(len(expected), 32767)
            report = importer.batch_export([source.parent], root / 'out')
            self.assertEqual(report['success_count'], 1)
            self.assertEqual(report['error_count'], 0)
            book = openpyxl.load_workbook(report['index'], read_only=True)
            try:
                self.assertEqual(book['批次索引']['A2'].value, '王府井')
                chunks = list(book['超长文本续表'].values)[1:]
                self.assertEqual(''.join(row[3] for row in chunks if row[1] == 'H2'), expected)
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
