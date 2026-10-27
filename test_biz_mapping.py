import json
import tempfile
import unittest
from pathlib import Path

import biz_mapping as bm
import excel_import as importer


def shop(name="测试店", floor="商场1F", category="咖啡"):
    return f"{name}\n★\n★\n4.5\n12条\n￥28/人\n{category}\n八宝街\n{floor}\n"


def snapshot(text):
    return {"documents": [{"name": "Page", "text": "输入商户名、地点、找优惠\n全部楼层\n全部分类\n" + text}],
            "primary_document_index": 0, "captured_at": "2026/10/21 15:13:25"}


class BizMappingTests(unittest.TestCase):
    def test_new_taxonomy_loads_without_whitelist_error(self):
        table = bm.read_csv_table(r"C:\nonexist" if False else self._csv("细分业态,标准业态\n咖啡,餐饮\n密室/剧本杀,娱乐\n宠物,配套\n"))
        self.assertEqual(table.groups, ["餐饮", "娱乐", "配套"])
        self.assertEqual(table.resolve("咖啡").group, "餐饮")
        self.assertEqual(table.resolve("宠物").group, "配套")

    def _csv(self, text, encoding="utf-8-sig"):
        self._temp = self._temp or tempfile.TemporaryDirectory()
        path = Path(self._temp.name) / f"表{len(list(Path(self._temp.name).iterdir()))}.csv"
        path.write_bytes(text.encode(encoding))
        return path

    def setUp(self):
        self._temp = None

    def tearDown(self):
        if self._temp:
            self._temp.cleanup()

    def test_duplicate_key_with_different_groups_is_explicit(self):
        table = bm.read_csv_table(self._csv("细分业态,标准业态\n其他,娱乐\n其他,配套\n咖啡,餐饮\n"))
        self.assertEqual(table.resolve("其他").method, "歧义")
        self.assertEqual(table.resolve("其他").group, "其他/待复核")
        self.assertTrue(any("同时对应" in w for w in table.warnings))

    def test_csv_encoding_and_delimiter_and_title_row(self):
        for encoding in ("gb18030", "utf-8-sig"):
            with self.subTest(encoding=encoding):
                table = bm.read_csv_table(self._csv("本地业态整理\r\n细分业态;标准业态;备注\r\n咖啡;餐饮;新版\r\n", encoding))
                self.assertEqual(table.resolve("咖啡").group, "餐饮")
                self.assertEqual(table.resolve("咖啡").method, "精确")

    def test_optional_alias_column_is_used(self):
        table = bm.read_csv_table(self._csv("细分业态,标准业态,别名\n茶饮果汁,餐饮;奶茶;鲜榨\n"))
        found = table.resolve("奶茶")
        self.assertEqual((found.category, found.group, found.method), ("茶饮果汁", "餐饮", "表内别名"))

    def test_xlsx_table_is_read_from_any_sheet(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "对照.xlsx"
            importer.write_xlsx(path, [("维护说明", [["项目", "说明"], ["x", "y"]], [20, 80]),
                                       ("业态对照表", [["细分业态", "标准业态"], ["火锅", "餐饮"]], [20, 20])])
            table = bm.read_table(path)
            self.assertEqual(table.resolve("火锅").group, "餐饮")
            self.assertEqual(table.source, path)
            self.assertTrue(table.fingerprint)

    def test_builtin_alias_and_keyword_fallback(self):
        table = bm.default_mapping()
        self.assertEqual((table.resolve("服装").category, table.resolve("服装").group), ("服饰", "零售"))
        self.assertEqual((table.resolve("川菜馆").category, table.resolve("川菜馆").group), ("川菜", "餐饮"))
        self.assertEqual((table.resolve("某某骨汤麻辣烫").category, table.resolve("某某骨汤麻辣烫").group),
                         ("小吃快餐/简餐", "餐饮"))
        self.assertEqual(table.resolve("骨汤麻辣烫").method, "关键词")
        self.assertEqual(table.resolve("智能按摩空间").group, "其他/待复核")
        self.assertEqual(table.resolve("智能按摩空间").method, "未匹配")

    def test_discover_prefers_newest_modified_table(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            old = root / "业态对照表.csv"
            old.write_text("细分业态,标准业态\n咖啡,餐饮\n", encoding="utf-8-sig")
            newest = root / "业态对照表_Excel参考.xlsx"
            importer.write_xlsx(newest, [("业态对照表", [["细分业态", "标准业态"], ["咖啡", "饮品"]], [20, 20])])
            import os
            os.utime(old, (1, 1))
            table, path, note, candidates = bm.discover_table([root])
            self.assertEqual(path, newest)
            self.assertEqual(table.resolve("咖啡").group, "饮品")
            self.assertIn("业态对照表.csv", note)

    def test_program_dir_table_wins_over_newer_data_dir_copy(self):
        """程序目录里的表是配置：数据目录里放了一份更新的同名表，也不抢。"""
        import os
        import time
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            app = root / "程序"
            app.mkdir()
            data = root / "采集数据" / "captures"
            data.mkdir(parents=True)
            program = app / "业态对照表.csv"
            program.write_text("细分业态,标准业态\n咖啡,餐饮\n", encoding="utf-8-sig")
            newer = root / "采集数据" / "业态对照表.csv"
            newer.write_text("细分业态,标准业态\n咖啡,饮品\n", encoding="utf-8-sig")
            stamp = time.time() + 100
            os.utime(newer, (stamp, stamp))
            table, path, note, candidates = bm.discover_table([app, data])
            self.assertEqual(path, program)
            self.assertEqual(table.resolve("咖啡").group, "餐饮")
            self.assertIn("业态对照表.csv", note)
            self.assertIn("业态对照表.csv", candidates[1].name)

    def test_discover_falls_back_to_builtin_and_reports(self):
        with tempfile.TemporaryDirectory() as temp:
            table, path, note, candidates = bm.discover_table([Path(temp)])
            self.assertIsNone(path)
            self.assertIsNone(table)
            self.assertIn("内置", note)

    def test_batch_export_uses_xlsx_mapping_and_marks_pending(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            table = root / "业态对照表.xlsx"
            importer.write_xlsx(table, [("业态对照表", [["细分业态", "标准业态"], ["咖啡", "餐饮"]], [20, 20])])
            folder = root / "captures" / "测试商场"
            folder.mkdir(parents=True)
            (folder / "capture.json").write_text(json.dumps(snapshot(shop() + shop("隔壁", "商场2F", "骨汤麻辣烫"))), encoding="utf-8")
            report = importer.batch_export([root / "captures"], root / "out", None, table)
            self.assertEqual(report["success_count"], 1)
            self.assertEqual(report["mapping"]["fingerprint"], bm.read_table(table).fingerprint)
            import openpyxl
            book = openpyxl.load_workbook(report["outputs"][0], read_only=True)
            try:
                rows = list(book["品牌明细"].values)[1:]
                self.assertEqual([row[6] for row in rows], ["餐饮", "其他/待复核"])
                self.assertEqual([row[7] for row in rows], ["咖啡", "骨汤麻辣烫"])
                self.assertIn("业态映射", book.sheetnames)
                self.assertIn("业态待补清单", book.sheetnames)
                self.assertEqual(book["业态待补清单"].max_row, 2)
            finally:
                book.close()
            index = openpyxl.load_workbook(report["index"], read_only=True)
            try:
                self.assertIn("业态待补清单", index.sheetnames)
                self.assertEqual(list(index["批次索引"].values)[1][8], "50.0%")
            finally:
                index.close()

    def test_batch_export_survives_broken_mapping_file(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            broken = root / "业态对照表.xlsx"
            broken.write_bytes(b"not an excel file")
            folder = root / "captures" / "测试商场"
            folder.mkdir(parents=True)
            (folder / "capture.json").write_text(json.dumps(snapshot(shop())), encoding="utf-8")
            report = importer.batch_export([root / "captures"], root / "out", None, broken)
            self.assertEqual(report["success_count"], 1)
            self.assertIn("跳过无效文件", report["mapping"]["note"])

    def test_normalize_keeps_slash_and_case_equivalent(self):
        self.assertEqual(bm.normalize("ＫＴＶ"), bm.normalize("ktv"))
        self.assertEqual(bm.normalize("小吃快餐／简餐"), bm.normalize("小吃快餐/简餐"))
        self.assertEqual(bm.normalize(" 咖啡  "), "咖啡")


if __name__ == "__main__":
    unittest.main()
